import frappe

def validate(doc, method=None):
    """Server-side: Ensure resume_attachment is a PDF file only."""
    if not doc.is_new():
        return

    if doc.get("resume_attachment"):
        url = doc.resume_attachment
        if not url.lower().endswith(".pdf"):
            frappe.throw(
                frappe._("Resume attachment must be a PDF file. Please upload a .pdf file."),
                title=frappe._("Invalid File Type")
            )

def create_resume_text_field():
    if not frappe.db.get_value("Custom Field", {"dt": "Job Applicant", "fieldname": "custom_resume_extraction"}):
        frappe.get_doc({
            "doctype": "Custom Field",
            "dt": "Job Applicant",
            "fieldname": "custom_resume_extraction",
            "label": "Resume Extraction",
            "fieldtype": "Text",
            "hidden": 1
        }).insert(ignore_permissions=True)

def extract_resume_text(doc=None, method=None, doc_name=None, doctype=None):
    # Determine the actual Document object and Doctype
    if type(doc) == str: # Backward compatibility if still called the old way
        doc = frappe.get_doc("Job Applicant", doc)
    elif not doc and doc_name and doctype:
        doc = frappe.get_doc(doctype, doc_name)
    
    current_doctype = doc.doctype if doc else doctype

    if current_doctype == "Job Applicant" and not hasattr(frappe.local, "resume_extraction_enabled"):
        create_resume_text_field()
        frappe.local.resume_extraction_enabled = True

    files = frappe.get_all("File", filters={"attached_to_doctype": current_doctype, "attached_to_name": doc.name})
    
    extracted_text = ""
    for f in files:
        file_doc = frappe.get_doc("File", f.name)
        text = extract_text_from_file(file_doc)
        if text:
            extracted_text += " " + text

    target_fieldname = "custom_resume_extraction" if current_doctype == "Job Applicant" else "resume_extraction"
            
    if extracted_text and doc.get(target_fieldname) != extracted_text:
        doc.db_set(target_fieldname, extracted_text.strip()[:65000])

def file_changed(doc, method=None):
    if doc.attached_to_doctype in ["Job Applicant", "Application Form"] and doc.attached_to_name:
        frappe.enqueue(
            "techner.techner.custom.job_applicant.extract_resume_text",
            doc_name=doc.attached_to_name,
            doctype=doc.attached_to_doctype,
            queue="short",
            timeout=300
        )

def extract_text_from_file(file_doc):
    try:
        import os
        file_path = file_doc.get_full_path()
        if not os.path.exists(file_path):
            return ""
            
        ext = file_doc.file_name.split(".")[-1].lower() if file_doc.file_name else ""
        text = ""
        
        if ext == "pdf":
            import pypdf
            reader = pypdf.PdfReader(file_path)
            for page in reader.pages:
                t = page.extract_text()
                if t: text += t + " "
                
        elif ext == "docx":
            import docx
            file_obj = docx.Document(file_path)
            for para in file_obj.paragraphs:
                text += para.text + " "
                
        elif ext in ["xlsx", "xls"]:
            import openpyxl
            wb = openpyxl.load_workbook(file_path, data_only=True)
            for sheet in wb.worksheets:
                for row in sheet.iter_rows(values_only=True):
                    row_texts = [str(cell) for cell in row if cell is not None]
                    if row_texts:
                        text += " ".join(row_texts) + " "
                        
        elif ext == "txt":
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                text = f.read()
        
        return text
    except Exception as e:
        frappe.log_error(frappe.get_traceback(), f"Resume Extraction Error for {file_doc.name}")
        return ""


@frappe.whitelist(allow_guest=True)
def bulk_extract_existing_attachments(doctype="Job Applicant"):
    """
    Sab existing Job Applicant records ki attachments extract karo
    jinki custom_resume_extraction field empty ho
    """
    records = frappe.get_all(
        doctype,
        filters={"custom_resume_extraction": ["in", ["", None]]},
        fields=["name"]
    )
    
    total = len(records)
    success = 0
    failed = 0
    
    for r in records:
        try:
            doc = frappe.get_doc(doctype, r.name)
            extract_resume_text(doc=doc)
            frappe.db.commit()
            success += 1
        except Exception as e:
            frappe.log_error(frappe.get_traceback(), f"Bulk Extraction Error: {r.name}")
            failed += 1
    
    return {
        "total": total,
        "success": success,
        "failed": failed
    }


@frappe.whitelist()
def extract_single_record(docname, doctype="Job Applicant"):
    """
    Single record ki attachments extract karo — force update bhi kare
    """
    doc = frappe.get_doc(doctype, docname)
    
    files = frappe.get_all(
        "File",
        filters={
            "attached_to_doctype": doctype,
            "attached_to_name": docname
        },
        fields=["name", "file_name"]
    )
    
    if not files:
        return {"status": "no_files", "message": "Koi attachment nahi mili"}
    
    extracted_text = ""
    extracted_files = []
    
    for f in files:
        file_doc = frappe.get_doc("File", f.name)
        text = extract_text_from_file(file_doc)
        if text:
            extracted_text += " " + text
            extracted_files.append(f.file_name)
    
    if extracted_text:
        doc.db_set("custom_resume_extraction", extracted_text.strip()[:65000])
        frappe.db.commit()
        return {
            "status": "success",
            "message": f"{len(extracted_files)} file(s) extract hui",
            "files": extracted_files
        }
    else:
        return {
            "status": "no_text",
            "message": "Files mili lekin text extract nahi hua"
        }